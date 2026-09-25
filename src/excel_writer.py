"""
excel_writer.py — Writes scored jobs into an Excel tracker, in the same
style as your existing job-search tracker, so you can review and click
through manually.
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

NAVY = "1F3864"
GREEN = "C6E0B4"
YELLOW = "FFF2CC"
RED = "F8CBAD"

HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=11)
HEADER_FILL = PatternFill("solid", start_color=NAVY)
NORMAL_FONT = Font(name="Arial", size=10)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

HEADERS = [
    "Match Score", "Company", "Job Title", "Location", "Source",
    "Matched Skills", "Gap Flags", "Posted", "Salary Range",
    "Application Link", "Cover Letter Draft", "Status", "Notes",
]


def write_results(jobs, out_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Job Agent Results"
    ws.sheet_view.showGridLines = False

    for i, h in enumerate(HEADERS, 1):
        ws.cell(row=1, column=i, value=h)
        ws.cell(row=1, column=i).font = HEADER_FONT
        ws.cell(row=1, column=i).fill = HEADER_FILL
        ws.cell(row=1, column=i).alignment = CENTER
        ws.cell(row=1, column=i).border = BORDER
    ws.row_dimensions[1].height = 28

    for r, job in enumerate(jobs, start=2):
        salary = ""
        if job.get("salary_min") or job.get("salary_max"):
            salary = f"{job.get('salary_min', '') or ''} - {job.get('salary_max', '') or ''}"
        row = [
            job["match_score"],
            job.get("company", ""),
            job.get("title", ""),
            job.get("location", ""),
            job.get("source", ""),
            ", ".join(job.get("matched_skills", [])),
            ", ".join(job.get("gap_flags", [])) or "-",
            job.get("posted", ""),
            salary,
            job.get("url", ""),
            job.get("cover_letter_path", "") or "-",
            "Not Applied",
            "",
        ]
        for c, val in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=val)
            cell.font = NORMAL_FONT
            cell.alignment = WRAP
            cell.border = BORDER
        ws.row_dimensions[r].height = 50

        score_cell = ws.cell(row=r, column=1)
        score_cell.alignment = CENTER
        score_cell.font = Font(name="Arial", bold=True, size=10)
        if job["match_score"] >= 75:
            score_cell.fill = PatternFill("solid", start_color=GREEN)
        elif job["match_score"] >= 55:
            score_cell.fill = PatternFill("solid", start_color=YELLOW)
        else:
            score_cell.fill = PatternFill("solid", start_color=RED)

    widths = [10, 22, 26, 18, 16, 30, 20, 14, 16, 34, 26, 14, 24]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = "A2"
    wb.save(out_path)
